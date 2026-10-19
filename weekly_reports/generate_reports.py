from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT.parent / "Project-I_Weekly_Plan_and_Evaluation_Framework - Student.docx"
OUTPUT = ROOT / "weekly_reports"
EVIDENCE = OUTPUT / "evidence"
REPOSITORY_URL = "https://github.com/rashiyaom/SentinelCore"
PROJECT_TITLE = "SentinelCore: A Network Detection and Incident Response Platform"
MEMBERS = "Om Rashiya (24cs084), Yash Patel (24cs082)"


WEEKS = [
    {
        "date": "20 July 2026",
        "planned": "Define the SGP problem statement, identify target users, study existing home-network security tools, and establish the project scope.",
        "completed": "Finalized SentinelCore as a self-hosted network detection and incident-response platform. Defined the primary workflow: discover devices, collect network alerts, correlate suspicious activity, investigate incidents, contain threats, and generate evidence reports.",
        "commits": "Total: 0 commits (research and planning week).\nRepository work had not started; requirements and module boundaries were documented before implementation.",
        "issues": "Created: Need a security platform understandable to non-specialist home/small-network users.\nClosed: Agreed on a modular, local-first platform rather than a cloud-only monitoring service.",
        "milestone": "~5% complete — on track. Problem definition, audience, objectives, and success criteria approved by the team.",
        "blockers": "The initial scope was too broad. We reduced it to a staged pipeline with clearly separated discovery, detection, investigation, response, and reporting capabilities.",
        "next": "Prepare the threat model, compare IDS technologies, and select the backend, database, queue, and frontend stack.",
        "research": "Compared Suricata, Snort, and Zeek for visibility, rule support, and structured event output. Suricata was selected because EVE JSON integrates cleanly with an application pipeline while retaining mature signature-based detection. We also reviewed the NIST incident-response lifecycle to shape the product flow.",
        "screenshot": "week01_hero.png",
        "caption": "Live evidence: the cinematic SecureHome/SentinelCore product introduction used to communicate the project identity.",
    },
    {
        "date": "27 July 2026",
        "planned": "Create a threat model, define functional/non-functional requirements, and design a safe deployment architecture.",
        "completed": "Documented trust boundaries for the browser, API, database, Redis queue, sensor, and privileged network helper. Defined RBAC roles, auditability, retention, least privilege, local deployment, and recoverable operations as core requirements.",
        "commits": "Total: 0 commits (architecture and research week).\nOutputs were the threat model, data-flow plan, and milestone breakdown used for later implementation.",
        "issues": "Created: Raw packet capture and firewall operations require privileges that the web API must not receive.\nClosed: Introduced a narrow Unix-socket helper boundary so the API can remain unprivileged.",
        "milestone": "~10% complete — on track. Architecture and security constraints were established before coding.",
        "blockers": "Balancing useful network visibility with safe privilege separation. The design isolates raw-socket operations and exposes only allow-listed helper commands.",
        "next": "Design the database entities, API contracts, Docker topology, and initial repository structure.",
        "research": "Studied container capability controls, Unix-domain socket permissions, secure refresh-token patterns, PostgreSQL transaction guarantees, and Redis-backed work coordination. The design adopted defense in depth: capability allowlists, unprivileged API execution, append-only audit records, and explicit role checks.",
        "screenshot": "week02_landing.png",
        "caption": "Live evidence: the public landing page explains discovery, detection, triage, and response in first-time-user language.",
    },
    {
        "date": "03 August 2026",
        "planned": "Translate the architecture into modules, database models, API routes, navigation, and a 12-week implementation plan.",
        "completed": "Split the platform into authentication, asset discovery, Suricata sensor, event pipeline, search, correlation, incidents, reports, firewall response, PCAP analysis, and threat intelligence. Defined FastAPI + SQLAlchemy/Alembic, PostgreSQL, Redis, React/Vite, Nginx, and Docker Compose as the implementation stack.",
        "commits": "Total: 0 commits (technical design week).\nPrepared module specifications M03–M12 and acceptance criteria for each milestone.",
        "issues": "Created: A single service would mix privileged capture, API traffic, and long-running ingestion.\nClosed: Separated API, worker, helper, frontend, proxy, PostgreSQL, and Redis responsibilities.",
        "milestone": "~18% complete — on track. Data model and module contracts were ready for implementation.",
        "blockers": "The event data volume could make offset pagination and one large table slow. The plan adopted monthly partitions and keyset pagination.",
        "next": "Initialize the repository, implement Docker services, health checks, migrations, and the first authentication screen.",
        "research": "Evaluated synchronous versus asynchronous API/database access, migration strategies, partitioned event storage, and React routing approaches. Async FastAPI and asyncpg were selected for concurrent event/search workloads, while Alembic provides reproducible schema evolution.",
        "screenshot": "week03_login.png",
        "caption": "Live evidence: the secure sign-in interface produced from the authentication and role-access design.",
    },
    {
        "date": "10 August 2026",
        "planned": "Create the repository and establish a reproducible full-stack development environment.",
        "completed": "Initialized SentinelCore, created backend/frontend/helper structures, added Docker Compose services, Nginx routing, PostgreSQL and Redis, environment configuration, and application health endpoints. Established the initial React shell and FastAPI service.",
        "commits": "Total: 2 commits.\n6a9c10c — “Initial commit”\n96396ba — “initial project setup”",
        "issues": "Created: Coordinating service startup and database readiness.\nClosed: Added health checks, dependency ordering, and a single Compose workflow for the entire stack.",
        "milestone": "~26% complete — on track. The full skeleton could be built and launched consistently.",
        "blockers": "Container networking and host-port conflicts required separating internal service names from loopback-only host mappings.",
        "next": "Complete container health, implement authentication/RBAC, seed an administrator, and add immutable auditing.",
        "research": "Reviewed Docker Compose health semantics, reverse-proxy routing, secret injection, and reproducible environment configuration. PostgreSQL is exposed only on loopback for tools, while application traffic stays on the private Compose network.",
        "screenshot": "week04_dashboard.png",
        "caption": "Live evidence: authenticated overview dashboard reached through the complete Docker/Nginx/API/frontend stack.",
    },
    {
        "date": "17 August 2026",
        "planned": "Stabilize the Docker stack and implement secure authentication, role-based access, session refresh, and audit logging.",
        "completed": "Completed the container stack health flow, added Argon2 password hashing, JWT access tokens with HttpOnly refresh cookies, login throttling, seeded admin support, protected React routes, role-aware navigation, and audit-log services.",
        "commits": "Total: 1 milestone commit.\n9fc44f2 — “M0: Docker stack complete, health endpoint working”\nAuthentication work continued in the later consolidated M1–M5 commit.",
        "issues": "Created: Access tokens must not persist in localStorage; repeated login attempts require throttling.\nClosed: Kept access tokens in memory, used rotating refresh cookies, and added server-side throttling and role checks.",
        "milestone": "~36% complete — on track. Secure access control and the deployment foundation were functional.",
        "blockers": "Maintaining a session without exposing durable tokens to JavaScript storage. The refresh endpoint and in-memory access token model resolved this.",
        "next": "Implement asset discovery, scan parsing, device inventory, and the privileged helper protocol.",
        "research": "Studied OWASP authentication guidance, Argon2 password hashing, short-lived bearer tokens, refresh-token cookies, RBAC enforcement, and tamper-evident audit design. Authorization is checked again in the backend even when navigation is hidden in the UI.",
        "screenshot": "week05_assets.png",
        "caption": "Live evidence: role-protected asset inventory and network discovery workspace.",
    },
    {
        "date": "24 August 2026",
        "planned": "Build asset discovery and begin integrating Suricata as the network sensor.",
        "completed": "Implemented allow-listed helper operations, Nmap scan execution/parsing, asset/port/scan models, scan history, validation, asset detail pages, Suricata configuration, rule sources, rule overrides, health metrics, and safe ruleset reload behavior.",
        "commits": "Total: 0 standalone commits this review period.\nThe feature branch was kept as a working milestone and later consolidated in d79a37c (“update the project to m5”).",
        "issues": "Created: Untrusted scan targets and rule text could cross the privilege boundary.\nClosed: Added strict network/argument validation, allow-listed helper operations, and verify-before-activate rule staging.",
        "milestone": "~45% complete — on track. Discovery and sensor administration were operational in development.",
        "blockers": "Nmap and Suricata require network access that differs across Linux and macOS Docker hosts. The helper owns capabilities while the API remains isolated.",
        "next": "Normalize EVE JSON, persist events reliably, track offsets, create partitions, and expose pipeline health.",
        "research": "Studied Nmap XML parsing, input validation for command execution, Suricata EVE JSON, rule update workflows, AF_PACKET capture, and atomic configuration activation. A staged ruleset is verified before replacing the active set to avoid sensor downtime.",
        "screenshot": "week06_sensor.png",
        "caption": "Live evidence: Suricata sensor health, throughput, rule sources, and override controls.",
    },
    {
        "date": "31 August 2026",
        "planned": "Create a resilient event-ingestion pipeline from Suricata EVE JSON to PostgreSQL.",
        "completed": "Built a dedicated worker, EVE tailer, event normalizer, batch writer, Redis-persisted read offsets, monthly PostgreSQL partitions, retention utilities, pipeline status endpoints, and event model validation. Kept ingestion outside the API process to prevent duplicate readers.",
        "commits": "Total: 0 standalone commits this review period.\nPipeline work was reviewed continuously and included in the consolidated M1–M5 milestone commit d79a37c.",
        "issues": "Created: Worker restarts could reread events or lose file position; multiple API replicas could duplicate ingestion.\nClosed: Moved tailing to one worker and persisted offsets in Redis with idempotent database writes.",
        "milestone": "~54% complete — on track. Detection events could flow reliably from sensor output into partitioned storage.",
        "blockers": "Handling log rotation, partial JSON lines, malformed events, and database backpressure without halting the pipeline.",
        "next": "Implement fast event search, facets, saved searches, and rule-driven correlation.",
        "research": "Compared file-tail reliability patterns, at-least-once delivery, checkpointing, batch insertion, and time-based partitioning. The worker acknowledges progress only after durable writes, and malformed lines are isolated rather than stopping ingestion.",
        "screenshot": "week07_events.png",
        "caption": "Live evidence: event explorer with time filters, severity/protocol controls, facets, and saved searches.",
    },
    {
        "date": "07 September 2026",
        "planned": "Add scalable event exploration and a correlation engine that can turn raw alerts into actionable findings.",
        "completed": "Designed keyset-paginated event queries, filter/facet contracts, saved searches, and correlation rules stored as data. Implemented evaluator patterns for threshold, sequence, rare-event, and beacon-like behavior, plus candidate creation and rule-run tracking.",
        "commits": "Total: 0 standalone commits this review period.\nSearch and correlation were integrated as feature milestones and committed after joint testing on 21 September.",
        "issues": "Created: Offset pagination degrades on large event tables; hard-coded detectors are difficult to tune.\nClosed: Adopted indexed keyset pagination and editable rules-as-data with bounded evaluator inputs.",
        "milestone": "~63% complete — on track. Analysts could search evidence and define repeatable detections.",
        "blockers": "Correlation needed deterministic windows and duplicate suppression so repeated worker passes would not create duplicate candidates.",
        "next": "Complete incident promotion, assignment, status/severity changes, evidence linking, and immutable history.",
        "research": "Studied keyset versus offset pagination, time-window correlation, threshold and sequence detection, beacon periodicity, and false-positive control. Rule state and runs are persisted so detections remain explainable and auditable.",
        "screenshot": "week08_correlation.png",
        "caption": "Live evidence: configurable correlation rule and candidate workspace.",
    },
    {
        "date": "14 September 2026",
        "planned": "Integrate M1–M5 work and implement the analyst incident lifecycle.",
        "completed": "Merged authentication, asset discovery, sensor, and event-pipeline foundations in one reviewed milestone. Added incident creation from candidates, human-readable numbering, status/severity/assignee changes, event attachment, comments, and append-only incident history.",
        "commits": "Total: 1 major consolidated commit.\nd79a37c — “update the project to m5” (130 files; authentication, assets, sensor, pipeline, tests, and UI foundation).",
        "issues": "Created: Candidates needed traceable promotion without losing original evidence.\nClosed: Incident records retain candidate/event relationships and record every lifecycle change in history.",
        "milestone": "~73% complete — on track. The platform supported discovery through investigation as one integrated workflow.",
        "blockers": "The milestone was large because related database migrations, worker logic, API contracts, tests, and frontend screens had to land together.",
        "next": "Finalize search/correlation commits, reporting, scheduled generation, and safe firewall containment.",
        "research": "Reviewed incident-response case management, chain of custody, analyst assignment, evidence preservation, and append-only change history. The model separates machine-generated candidates from analyst-managed incidents while maintaining the source trail.",
        "screenshot": "week09_incidents.png",
        "caption": "Live evidence: incident queue supporting triage, assignment, status, severity, and evidence review.",
    },
    {
        "date": "21 September 2026",
        "planned": "Complete search, correlation, and incident management; start evidence reporting and response actions.",
        "completed": "Completed event search with facets/saved searches, four correlation evaluator types and live worker execution, and incident lifecycle management. Defined asynchronous report generation and a protected firewall response model with expiry and recovery behavior.",
        "commits": "Total: 3 feature commits.\n567e05d — “M6: event search — keyset-paginated query API, facets, saved searches”\n76d671a — “M7: correlation engine — rules-as-data, four evaluators, live worker task”\nba4c083 — “M8: incident management — lifecycle, promotion from candidates, history”",
        "issues": "Created: Long report jobs should not block API requests; containment must never block protected addresses.\nClosed: Designed queued report generation and mandatory protection/TTL checks before firewall changes.",
        "milestone": "~82% complete — on track. Detection-to-incident workflow was complete and response/reporting work had begun.",
        "blockers": "Keeping evaluator queries bounded and ensuring incident promotion remained idempotent under repeated worker execution.",
        "next": "Finish PDF/CSV/JSON reports, schedules, firewall containment, PCAP upload/analysis, and threat intelligence.",
        "research": "Studied asynchronous report pipelines, reproducible evidence exports, iptables containment, protected-IP guardrails, TTL-based rollback, and crash reconciliation. Response actions require explicit user intent and retain a complete audit trail.",
        "screenshot": "week10_reports.png",
        "caption": "Live evidence: report generation and scheduling workspace for PDF, CSV, and JSON evidence.",
    },
    {
        "date": "27 September 2026",
        "planned": "Complete reporting, containment, PCAP analysis, and threat-intelligence enrichment.",
        "completed": "Implemented four asynchronous report types, report schedules, PDF/CSV/JSON outputs, protected and TTL-bounded firewall blocks, PCAP upload/parser isolation, flow/artifact extraction, IOC/source management, feed normalization, and event/incident enrichment.",
        "commits": "Total: 5 milestone commits.\n87f7e20 — “M9: reporting backend — four report types, async PDF/CSV/JSON generation, scheduling”\ne536b79 — “M9: reporting frontend — reports table, new-report form, admin schedules”\nfe1ca48 — “M10: firewall containment — protection guard, iptables ops, TTL-bounded blocks”\n7082a69 — “M5 to M10 work”\n0c98a3e — “complte all modules” (M11/M12 integration and E2E coverage)",
        "issues": "Created: Uploaded captures and external feeds are untrusted; firewall operations are high impact.\nClosed: Isolated parsing, bounded file/resource use, normalized IOCs, protected critical addresses, and enforced automatic unblock expiry.",
        "milestone": "~91% complete — on track. All planned functional modules were integrated end to end.",
        "blockers": "Safely parsing large/malformed captures and feeds while keeping memory, database work, and external requests bounded.",
        "next": "Run end-to-end validation, harden security controls, improve documentation/UI clarity, and prepare the final demonstration.",
        "research": "Reviewed PCAP parser isolation, MIME/size validation, flow reconstruction, IOC normalization (IP/domain/hash), safe feed retrieval, SSRF controls, and reversible containment. External indicators remain attributable to their source and match evidence.",
        "screenshot": "week11_firewall.png",
        "caption": "Live evidence: guarded containment console with active blocks, history, protected targets, and expiry controls.",
    },
    {
        "date": "30 September 2026",
        "planned": "Validate the complete platform, fix integration defects, polish documentation, and prepare the faculty demonstration.",
        "completed": "Ran backend, helper, pipeline, and browser-level checks; verified authentication, discovery, detection, search, incidents, reports, firewall, PCAP, and intelligence flows. Improved the README, test evidence, dashboard clarity, navigation, empty states, visual hierarchy, responsive behavior, and first-time-user guidance.",
        "commits": "Total: 2 completion commits.\n6945e77 — “updated and tested”\n19e4da8 — “Polish README with badges, TOC, and visual section dividers”\nLater hardening commits continue transparently in the linked repository history.",
        "issues": "Created: Integration testing exposed confusing empty states and several edge cases around parser output, reporting, and request security.\nClosed: Added clearer guidance, stronger validation, reproducible test documentation, and follow-up hardening tasks.",
        "milestone": "100% of the planned SGP scope demonstrated. Core platform complete; subsequent commits are maintenance, security hardening, and UI refinement.",
        "blockers": "Docker networking/capture behavior varies by host OS, and realistic packet data is environment-dependent. The demonstration uses health checks, bounded test fixtures, and explicit setup guidance.",
        "next": "Collect faculty feedback, continue security/performance hardening, expand realistic datasets, and prepare deployment documentation and final presentation material.",
        "research": "Applied OWASP guidance for SSRF, spoofed client IPs, rate limits, request validation, race prevention, and secure defaults. Also reviewed dashboard information hierarchy and progressive disclosure so a first-time user can understand system health and the next action immediately.",
        "screenshot": "week12_intelligence.png",
        "caption": "Live evidence: final threat-intelligence workspace for IOCs, matches, sources, and analyst filtering.",
    },
]


def set_cell_text(cell, text, size=9, bold=False):
    cell.text = ""
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        run = paragraph.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.bold = bold


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "5B3FD3")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def move_before_table(paragraph, table):
    table._tbl.addprevious(paragraph._p)


def add_evidence_section(document, faculty_table, week_number, data):
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run("Research & Technical Study")
    heading_run.bold = True
    heading_run.font.name = "Arial"
    heading_run.font.size = Pt(11)
    heading_run.font.color.rgb = RGBColor(42, 31, 78)
    move_before_table(heading, faculty_table)

    research = document.add_paragraph()
    research.paragraph_format.space_after = Pt(6)
    research.paragraph_format.line_spacing = 1.0
    research_run = research.add_run(data["research"])
    research_run.font.name = "Arial"
    research_run.font.size = Pt(9)
    move_before_table(research, faculty_table)

    evidence_heading = document.add_paragraph()
    evidence_heading.paragraph_format.space_before = Pt(3)
    evidence_heading.paragraph_format.space_after = Pt(3)
    evidence_heading.paragraph_format.keep_with_next = True
    evidence_run = evidence_heading.add_run(f"Weekly Evidence — Week {week_number}")
    evidence_run.bold = True
    evidence_run.font.name = "Arial"
    evidence_run.font.size = Pt(11)
    evidence_run.font.color.rgb = RGBColor(42, 31, 78)
    move_before_table(evidence_heading, faculty_table)

    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(3)
    image_run = image_paragraph.add_run()
    image_run.add_picture(str(EVIDENCE / data["screenshot"]), width=Inches(6.85))
    move_before_table(image_paragraph, faculty_table)

    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(4)
    caption_run = caption.add_run(data["caption"])
    caption_run.italic = True
    caption_run.font.name = "Arial"
    caption_run.font.size = Pt(8)
    move_before_table(caption, faculty_table)

    repository = document.add_paragraph()
    repository.paragraph_format.space_after = Pt(7)
    label = repository.add_run("GitHub Repository: ")
    label.bold = True
    label.font.name = "Arial"
    label.font.size = Pt(9)
    add_hyperlink(repository, REPOSITORY_URL, REPOSITORY_URL)
    move_before_table(repository, faculty_table)


def generate_report(week_number, data):
    document = Document(TEMPLATE)
    document.core_properties.title = f"24cs084 Week {week_number} Progress Report"
    document.core_properties.subject = PROJECT_TITLE
    document.core_properties.author = "Om Rashiya and Yash Patel"

    identity, progress, faculty = document.tables
    set_cell_text(identity.cell(0, 1), "CS26084", size=9, bold=True)
    set_cell_text(identity.cell(0, 3), f"Week {week_number}\n{data['date']}", size=9, bold=True)
    set_cell_text(identity.cell(1, 1), PROJECT_TITLE, size=9, bold=True)
    set_cell_text(identity.cell(2, 1), MEMBERS, size=9)

    fields = ["planned", "completed", "commits", "issues", "milestone", "blockers", "next"]
    for row_index, field in enumerate(fields, start=1):
        set_cell_text(progress.cell(row_index, 1), data[field], size=8.5)

    # Faculty evaluation must remain entirely blank for faculty use.
    for cell in faculty.rows[1].cells:
        set_cell_text(cell, "", size=9)

    add_evidence_section(document, faculty, week_number, data)

    # Keep the evaluation label with its empty row whenever Word paginates.
    for cell in faculty.rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True

    output_path = OUTPUT / f"24cs084_Week{week_number}.docx"
    document.save(output_path)
    return output_path


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for week_number, data in enumerate(WEEKS, start=1):
        path = generate_report(week_number, data)
        print(path.name)


if __name__ == "__main__":
    main()
